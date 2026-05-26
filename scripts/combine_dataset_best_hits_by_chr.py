#!/usr/bin/env python3

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


HIGHLIGHT_FILL = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")


def parse_args():
    parser = argparse.ArgumentParser(description="Combine per-chromosome dataset best-hit tables")
    parser.add_argument("--inputs", nargs="+", required=True, help="Per-chromosome best-hit TSV files")
    parser.add_argument("--output-tsv", required=True, help="Combined TSV output")
    parser.add_argument("--output-xlsx", required=True, help="Combined XLSX output")
    return parser.parse_args()


def write_excel_with_highlight(df, xlsx_path):
    p_cols = [col for col in df.columns if col.endswith("__top_p")]
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="dataset_best_hits")

    workbook = load_workbook(xlsx_path)
    sheet = workbook["dataset_best_hits"]
    col_idx = {cell.value: cell.column for cell in sheet[1]}
    for col_name in p_cols:
        idx = col_idx.get(col_name)
        if idx is None:
            continue
        for row in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row, column=idx).value
            try:
                if value is not None and float(value) < 5e-8:
                    sheet.cell(row=row, column=idx).fill = HIGHLIGHT_FILL
            except (TypeError, ValueError):
                continue

    workbook.save(xlsx_path)


def main():
    args = parse_args()
    frames = []
    for path in args.inputs:
        input_path = Path(path)
        if not input_path.exists():
            continue
        df = pd.read_csv(input_path, sep="\t", low_memory=False)
        if not df.empty:
            frames.append(df)

    if frames:
        out = pd.concat(frames, ignore_index=True)
        if {"chr", "start"}.issubset(out.columns):
            out["chr_sort"] = pd.to_numeric(out["chr"], errors="coerce")
            out["start_sort"] = pd.to_numeric(out["start"], errors="coerce")
            out = out.sort_values(["chr_sort", "start_sort"], na_position="last").drop(columns=["chr_sort", "start_sort"])
    else:
        out = pd.DataFrame()

    Path(args.output_tsv).parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(args.output_tsv, sep="\t", index=False)
    write_excel_with_highlight(out, args.output_xlsx)
    print(f"TSV written: {args.output_tsv}")
    print(f"XLSX written: {args.output_xlsx}")


if __name__ == "__main__":
    main()