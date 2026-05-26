#!/usr/bin/env python3

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


HIGHLIGHT_FILL = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Create a simple per-merged-locus table with the best SNP and best p-value "
            "for each dataset using Step 3 chromosome-level subset GWAS tables."
        )
    )
    parser.add_argument("--config", required=True, help="Path to analysis.yml")
    parser.add_argument("--summary-tsv", required=True, help="Merged Step 7 summary TSV")
    parser.add_argument("--output-tsv", required=True, help="Output simple TSV")
    parser.add_argument("--output-xlsx", help="Output simple XLSX")
    parser.add_argument("--skip-xlsx", action="store_true", help="Skip writing XLSX output")
    return parser.parse_args()


def read_yaml(path):
    with open(path, "r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def find_first(columns, candidates):
    lower_map = {str(col).lower(): str(col) for col in columns}
    for candidate in candidates:
        match = lower_map.get(candidate.lower())
        if match is not None:
            return match
    return None


def normalize_chr(value):
    if pd.isna(value):
        return np.nan
    text = str(value).strip()
    if text.lower().startswith("chr"):
        text = text[3:]
    return text


def detect_columns(path):
    if not Path(path).exists():
        return None

    header = pd.read_csv(path, sep="\t", nrows=0)
    columns = list(header.columns)
    chr_col = find_first(columns, ["CHROM", "CHR", "chrom", "chr", "chromosome"])
    pos_col = find_first(columns, ["POS", "pos", "BP", "bp", "position"])
    p_col = find_first(
        columns,
        [
            "p",
            "P",
            "pvalue",
            "P_VALUE",
            "p_local_ancestry",
            "p_dosage_i",
            "p_dosage_j",
        ],
    )
    snp_col = find_first(columns, ["ID", "id", "rsid", "snpid", "markername", "varid"])
    if chr_col is None or pos_col is None or p_col is None:
        return None
    return chr_col, pos_col, p_col, snp_col


def normalize_chr_token(value):
    normalized = normalize_chr(value)
    if pd.isna(normalized):
        return None
    text = str(normalized).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def get_subset_path(results_dir, dataset, table, chrom):
    chrom_token = normalize_chr_token(chrom)
    if chrom_token is None:
        return None
    return (
        Path(results_dir)
        / "gwas_subset"
        / dataset
        / f"{table}_subset_dir"
        / f"{dataset}_{table}_chr{chrom_token}_subset.tsv"
    )


def load_subset_table(path):
    detected = detect_columns(path)
    if detected is None:
        return pd.DataFrame(columns=["CHR", "POS", "SNP", "PVAL"])

    chr_col, pos_col, p_col, snp_col = detected
    usecols = [chr_col, pos_col, p_col] + ([snp_col] if snp_col else [])
    df = pd.read_csv(path, sep="\t", usecols=usecols, low_memory=False)
    if df.empty:
        return pd.DataFrame(columns=["CHR", "POS", "SNP", "PVAL"])

    rename_map = {chr_col: "CHR", pos_col: "POS", p_col: "PVAL"}
    if snp_col:
        rename_map[snp_col] = "SNP"
    df = df.rename(columns=rename_map)
    if "SNP" not in df.columns:
        df["SNP"] = np.nan

    df["CHR"] = df["CHR"].map(normalize_chr)
    df["POS"] = pd.to_numeric(df["POS"], errors="coerce")
    df["PVAL"] = pd.to_numeric(df["PVAL"], errors="coerce")
    df = df.dropna(subset=["CHR", "POS", "PVAL"])
    df = df[(df["PVAL"] > 0) & (df["PVAL"] <= 1)].copy()
    if df.empty:
        return pd.DataFrame(columns=["CHR", "POS", "SNP", "PVAL"])

    df["POS"] = df["POS"].astype(int)
    df["SNP"] = df["SNP"].astype(str)
    return df[["CHR", "POS", "SNP", "PVAL"]].sort_values(["POS", "PVAL"]).reset_index(drop=True)


def stream_best_hits_for_dataset(summary, dataset_entry, results_dir):
    dataset = str(dataset_entry["name"])
    best_p = np.full(len(summary), np.nan)
    best_snp = np.full(len(summary), np.nan, dtype=object)
    best_table = np.full(len(summary), np.nan, dtype=object)

    loci_by_chr = {}
    for row_idx, row in summary.iterrows():
        chr_value = normalize_chr_token(row["chr"])
        if chr_value is None:
            continue
        loci_by_chr.setdefault(chr_value, []).append((row_idx, int(row["start"]), int(row["end"])))

    for table_entry in dataset_entry.get("tables", []):
        table = str(table_entry["table_name"])
        for chr_value, loci in loci_by_chr.items():
            subset_path = get_subset_path(results_dir, dataset, table, chr_value)
            if subset_path is None or not subset_path.exists():
                continue
            subset_df = load_subset_table(subset_path)
            if subset_df.empty:
                continue

            for row_idx, locus_start, locus_end in loci:
                in_locus = subset_df[(subset_df["POS"] >= locus_start) & (subset_df["POS"] <= locus_end)]
                if in_locus.empty:
                    continue

                min_idx = in_locus["PVAL"].idxmin()
                min_row = in_locus.loc[min_idx]
                cur_best = best_p[row_idx]
                new_best = float(min_row["PVAL"])
                if pd.isna(cur_best) or new_best < cur_best:
                    best_p[row_idx] = new_best
                    best_snp[row_idx] = min_row["SNP"]
                    best_table[row_idx] = table

    return best_snp, best_p, best_table


def build_dataset_best_hits(summary_df, cfg):
    results_dir = cfg.get("results_dir", "results")
    summary = summary_df.copy()
    summary["chr"] = summary["chr"].map(normalize_chr)
    summary["start"] = pd.to_numeric(summary["start"], errors="coerce").astype(int)
    summary["end"] = pd.to_numeric(summary["end"], errors="coerce").astype(int)

    base_cols = [
        "locus_code",
        "chr",
        "start",
        "end",
        "length_bp",
        "length_kbp",
        "nearest_gene",
        "gene_distance_bp",
    ]
    output = summary[base_cols].copy()

    for dataset_entry in cfg.get("gwastables", []):
        dataset = str(dataset_entry["name"])
        best_snps, best_pvals, best_tables = stream_best_hits_for_dataset(summary, dataset_entry, results_dir)

        output[f"{dataset}__top_snp"] = best_snps
        output[f"{dataset}__top_p"] = best_pvals
        output[f"{dataset}__top_table"] = best_tables

    return output


def write_excel_with_highlight(df, xlsx_path):
    p_cols = [col for col in df.columns if col.endswith("__top_p")]
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="dataset_best_hits")

    workbook = load_workbook(xlsx_path)
    sheet = workbook["dataset_best_hits"]
    col_idx = {cell.value: cell.column for cell in sheet[1]}

    for col_name in p_cols:
        idx = col_idx.get(col_name)
        if idx is None:
            continue
        for row in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row, column=idx).value
            try:
                if value is not None and float(value) < 5e-8:
                    sheet.cell(row=row, column=idx).fill = HIGHLIGHT_FILL
            except (TypeError, ValueError):
                continue

    workbook.save(xlsx_path)


def main():
    args = parse_args()
    cfg = read_yaml(args.config)
    summary_df = pd.read_csv(args.summary_tsv, sep="\t", low_memory=False)
    out_df = build_dataset_best_hits(summary_df, cfg)

    Path(args.output_tsv).parent.mkdir(parents=True, exist_ok=True)
    out_df.to_csv(args.output_tsv, sep="\t", index=False)
    if not args.skip_xlsx:
        if not args.output_xlsx:
            raise ValueError("--output-xlsx is required unless --skip-xlsx is used")
        write_excel_with_highlight(out_df, args.output_xlsx)

    print(f"TSV written: {args.output_tsv}")
    if not args.skip_xlsx:
        print(f"XLSX written: {args.output_xlsx}")
    print(f"Rows: {len(out_df)}  Cols: {len(out_df.columns)}")


if __name__ == "__main__":
    main()