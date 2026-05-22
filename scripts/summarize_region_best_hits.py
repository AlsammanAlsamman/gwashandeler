#!/usr/bin/env python3
"""Summarize best SNP hit per configured region across all dataset/table subsets."""

from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import yaml

LIGHT_RED_FILL = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
GENOMEWIDE_THRESHOLD = 5e-8
SUGGESTIVE_THRESHOLD = 5e-5
SCIENTIFIC_FORMAT = "0.00E+00"
P_PATTERNS = (r"p_", r"pval", r"pvalue", r"^p$")
ID_PATTERNS = (r"^rsid$", r"^id$", r"^snpid$", r"^snp$", r"^markername$", r"^varid$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize best p-value and SNP id across region subsets")
    parser.add_argument("--config", required=True)
    parser.add_argument("--regions", required=True)
    parser.add_argument("--subsets-root", required=True)
    parser.add_argument("--output-tsv", required=True)
    parser.add_argument("--output-xlsx", required=True)
    return parser.parse_args()


def norm_chr(value: str) -> str:
    return re.sub(r"^chr", "", str(value).strip(), flags=re.IGNORECASE)


def sanitize_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value).strip())


def find_col(columns: list[str], patterns: tuple[str, ...], exclude: set[str] | None = None) -> str | None:
    exclude = exclude or set()
    for pattern in patterns:
        for col in columns:
            col_norm = str(col).strip().lower()
            if col_norm in exclude:
                continue
            if re.search(pattern, col_norm, re.IGNORECASE):
                return str(col)
    return None


def load_config(config_path: Path) -> list[dict[str, object]]:
    with config_path.open("r", encoding="utf-8") as fh:
        cfg = yaml.safe_load(fh)
    return list(cfg.get("gwastables", []))


def load_regions(regions_path: Path) -> list[dict[str, object]]:
    regions: list[dict[str, object]] = []
    with regions_path.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        for row in reader:
            locus = str(row.get("locus", "")).strip()
            chr_raw = str(row.get("chr", "")).strip()
            start_raw = str(row.get("start", "")).strip()
            end_raw = str(row.get("end", "")).strip()
            if not locus or not chr_raw or not start_raw or not end_raw:
                continue
            try:
                start = int(float(start_raw))
                end = int(float(end_raw))
            except ValueError:
                continue
            if end < start:
                start, end = end, start
            chr_norm = norm_chr(chr_raw)
            regions.append(
                {
                    "locus": locus,
                    "chr": chr_norm,
                    "start": start,
                    "end": end,
                    "region_id": sanitize_name(f"{locus}_chr{chr_norm}_{start}_{end}"),
                }
            )
    return regions


def summarize_subset(subset_path: Path) -> tuple[float | None, str | None]:
    if not subset_path.exists():
        return None, None

    df = pd.read_csv(subset_path, sep="\t", low_memory=False)
    if df.empty:
        return None, None

    p_col = find_col(list(df.columns), P_PATTERNS, {"pos", "position", "chrom", "chr", "bp", "allele"})
    id_col = find_col(list(df.columns), ID_PATTERNS)
    if p_col is None or id_col is None:
        return None, None

    df[p_col] = pd.to_numeric(df[p_col], errors="coerce")
    valid = df[(df[p_col] > 0) & (df[p_col] <= 1)].copy()
    if valid.empty:
        return None, None

    best_idx = valid[p_col].idxmin()
    best_p = float(valid.loc[best_idx, p_col])
    best_id = valid.loc[best_idx, id_col]
    if pd.isna(best_id):
        best_id = None
    else:
        best_id = str(best_id)
    return best_p, best_id


def write_excel_with_highlight(df: pd.DataFrame, xlsx_path: str, highlight_cols: list[str]) -> None:
    Path(os.path.dirname(xlsx_path)).mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="region_best_hits")

    wb = load_workbook(xlsx_path)
    ws = wb["region_best_hits"]
    col_idx = {cell.value: cell.column for cell in ws[1]}
    for col_name in highlight_cols:
        idx = col_idx.get(col_name)
        if idx is None:
            continue
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=idx)
            value = cell.value
            try:
                if value is None:
                    continue
                numeric_value = float(value)
                cell.number_format = SCIENTIFIC_FORMAT
                if numeric_value < GENOMEWIDE_THRESHOLD:
                    cell.fill = LIGHT_RED_FILL
                elif GENOMEWIDE_THRESHOLD <= numeric_value <= SUGGESTIVE_THRESHOLD:
                    cell.fill = YELLOW_FILL
            except (ValueError, TypeError):
                continue
    wb.save(xlsx_path)


def main() -> None:
    args = parse_args()
    subsets_root = Path(args.subsets_root)
    config_path = Path(args.config)
    regions_path = Path(args.regions)

    gwastables = load_config(config_path)
    regions = load_regions(regions_path)

    rows: list[dict[str, object]] = []
    for region in regions:
        row: dict[str, object] = {
            "locus": region["locus"],
            "chr": region["chr"],
            "start": region["start"],
            "end": region["end"],
        }
        region_id = str(region["region_id"])

        for dataset in gwastables:
            dataset_name = str(dataset["name"])
            for table in dataset.get("tables", []):
                table_name = str(table["table_name"])
                subset_path = subsets_root / dataset_name / f"{region_id}__{table_name}.tsv"
                best_p, best_id = summarize_subset(subset_path)
                prefix = f"{dataset_name}__{table_name}"
                row[f"{prefix}__best_p"] = best_p if best_p is not None else np.nan
                row[f"{prefix}__best_rsid"] = best_id

        rows.append(row)

    out = pd.DataFrame(rows)
    highlight_cols = [col for col in out.columns if str(col).endswith("__best_p")]

    output_tsv = Path(args.output_tsv)
    output_tsv.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(output_tsv, sep="\t", index=False)
    write_excel_with_highlight(out, args.output_xlsx, highlight_cols)


if __name__ == "__main__":
    main()
